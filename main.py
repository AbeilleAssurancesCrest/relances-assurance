from datetime import datetime
from io import BytesIO
import unicodedata

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from database import (init_db, add_client_avec_contrats, get_all_clients, 
                      log_relance, get_historique, delete_client, add_contrats_a_client,
                      update_contrat_details, update_commentaire, update_coordonnees,
                      get_rendez_vous_commerciaux, add_rendez_vous_commercial,
                      update_rendez_vous_commercial, delete_rendez_vous_commercial)

app = FastAPI()
init_db()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory="."), name="static")
templates = Jinja2Templates(directory="templates")

class ContratSchema(BaseModel):
    num_contrat: str
    type_vehicule: str = "Voiture"
    date_effet: str
    marque: str = ""
    immat: str = ""
    pieces: list[str] = []

class ClientSchema(BaseModel):
    nom: str
    prenom: str
    email: str = ""
    telephone: str = ""
    commentaire: str = ""
    contrats: list[ContratSchema]

class AjouterContratsSchema(BaseModel):
    contrats: list[ContratSchema]

class RelanceSchema(BaseModel):
    client_id: int
    email: str
    pieces: list[str]

class ContratUpdateSchema(BaseModel):
    contrat_id: int
    immat: str = ""
    pieces: list[str] = []

class CommentaireUpdateSchema(BaseModel):
    client_id: int
    commentaire: str

class CoordonneesUpdateSchema(BaseModel):
    client_id: int
    email: str = ""
    telephone: str = ""

class RendezVousCommercialSchema(BaseModel):
    client_id: int | None = None
    nom: str
    prenom: str = ""
    email: str = ""
    telephone: str = ""
    profil: str
    date_rdv: str
    besoin: str = ""
    statut: str = "RDV à venir"
    prochaine_action: str = ""
    commentaire: str = ""

@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    return templates.TemplateResponse(request=request, name="index.html")

@app.get("/api/dossiers")
def fetch_dossiers():
    return get_all_clients()

@app.get("/api/rendez-vous-commerciaux")
def fetch_rendez_vous_commerciaux():
    return get_rendez_vous_commerciaux()

def valider_rendez_vous_commercial(data: RendezVousCommercialSchema):
    if not data.nom.strip():
        raise HTTPException(status_code=400, detail="Le nom du contact est obligatoire.")
    if data.profil not in {"prospect_entrant", "prospect_demarche", "client_mono"}:
        raise HTTPException(status_code=400, detail="Choisis une catégorie valide.")
    if data.statut not in {"RDV à venir", "RDV réalisé", "Offre à préparer", "Offre envoyée", "À recontacter", "Gagné", "Perdu"}:
        raise HTTPException(status_code=400, detail="Choisis un statut valide.")
    try:
        datetime.fromisoformat(data.date_rdv.replace("Z", "+00:00"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Indique une date et une heure de rendez-vous valides.")
    if data.client_id is not None:
        client = next((client for client in get_all_clients() if client.get("id") == data.client_id), None)
        if client is None:
            raise HTTPException(status_code=404, detail="La fiche client sélectionnée est introuvable.")
        if data.profil != "client_mono" or len(client.get("contrats", [])) != 1:
            raise HTTPException(status_code=400, detail="Le lien client est réservé aux fiches avec exactement un contrat.")

@app.post("/api/rendez-vous-commerciaux")
def create_rendez_vous_commercial(data: RendezVousCommercialSchema):
    valider_rendez_vous_commercial(data)
    enregistre, detail = add_rendez_vous_commercial(data.dict())
    if not enregistre:
        raise HTTPException(status_code=502, detail=detail)
    return {"status": "ok"}

@app.put("/api/rendez-vous-commerciaux/{rendez_vous_id}")
def modify_rendez_vous_commercial(rendez_vous_id: int, data: RendezVousCommercialSchema):
    valider_rendez_vous_commercial(data)
    ok, detail = update_rendez_vous_commercial(rendez_vous_id, data.dict())
    if not ok:
        raise HTTPException(status_code=502, detail=detail)
    return {"status": "ok"}

@app.delete("/api/rendez-vous-commerciaux/{rendez_vous_id}")
def remove_rendez_vous_commercial(rendez_vous_id: int):
    ok, detail = delete_rendez_vous_commercial(rendez_vous_id)
    if not ok:
        raise HTTPException(status_code=502, detail=detail)
    return {"status": "ok"}

def normaliser_identite(value: str) -> str:
    sans_accents = "".join(
        caractere for caractere in unicodedata.normalize("NFKD", value or "")
        if not unicodedata.combining(caractere)
    )
    return " ".join(sans_accents.split()).casefold()

def verifier_numeros_contrat(contrats, clients_existants):
    numeros_saisis = set()
    proprietaires = {}
    for client in clients_existants:
        for contrat in client.get("contrats", []):
            numero = str(contrat.get("num_contrat", "")).strip()
            if numero:
                proprietaires.setdefault(numero, client)

    for contrat in contrats:
        numero = str(contrat.num_contrat).strip()
        if len(numero) != 8 or not numero.isdigit():
            raise HTTPException(status_code=400, detail="Le numéro de contrat doit comporter exactement 8 chiffres.")
        if numero in numeros_saisis:
            raise HTTPException(status_code=409, detail=f"Le numéro de contrat {numero} apparaît plusieurs fois dans le formulaire.")
        if numero in proprietaires:
            client = proprietaires[numero]
            nom_client = f"{client.get('prenom', '')} {client.get('nom', '')}".strip()
            raise HTTPException(status_code=409, detail=f"Le contrat {numero} est déjà enregistré pour {nom_client}.")
        numeros_saisis.add(numero)

@app.post("/api/dossiers")
def create_dossier(data: ClientSchema):
    if not data.email.strip() and not data.telephone.strip():
        raise HTTPException(status_code=400, detail="Ajoute au moins un moyen de contact : e-mail ou téléphone.")
    clients_existants = get_all_clients()
    identite_saisie = (normaliser_identite(data.nom), normaliser_identite(data.prenom))
    doublon_client = next((
        client for client in clients_existants
        if (normaliser_identite(client.get("nom", "")), normaliser_identite(client.get("prenom", ""))) == identite_saisie
    ), None)
    if doublon_client:
        raise HTTPException(status_code=409, detail="Ce client existe déjà. Ajoute le contrat à sa fiche existante.")
    verifier_numeros_contrat(data.contrats, clients_existants)
    contrats_list = [c.dict() for c in data.contrats]
    for contrat in contrats_list:
        contrat["immat"] = (contrat.get("immat") or "").strip().upper()
    add_client_avec_contrats(data.nom.strip().upper(), data.prenom, data.email, data.telephone, data.commentaire, contrats_list)
    return {"status": "ok"}

@app.post("/api/dossiers/{client_id}/contrats")
def add_contracts_to_existing_client(client_id: int, data: AjouterContratsSchema):
    clients_existants = get_all_clients()
    if not any(client.get("id") == client_id for client in clients_existants):
        raise HTTPException(status_code=404, detail="Client introuvable.")
    if not data.contrats:
        raise HTTPException(status_code=400, detail="Ajoute au moins un contrat.")
    verifier_numeros_contrat(data.contrats, clients_existants)
    contrats_list = [c.dict() for c in data.contrats]
    for contrat in contrats_list:
        contrat["immat"] = (contrat.get("immat") or "").strip().upper()
    if not add_contrats_a_client(client_id, contrats_list):
        raise HTTPException(status_code=502, detail="Impossible d'ajouter le contrat au dossier.")
    return {"status": "ok"}

@app.post("/api/contrats/update")
def modify_contrat(data: ContratUpdateSchema):
    ok, detail = update_contrat_details(data.contrat_id, data.immat.strip().upper(), data.pieces)
    if not ok:
        raise HTTPException(status_code=502, detail=detail)
    return {"status": "ok"}

@app.post("/api/dossiers/update_commentaire")
def modify_commentaire(data: CommentaireUpdateSchema):
    ok, detail = update_commentaire(data.client_id, data.commentaire)
    if not ok:
        raise HTTPException(status_code=502, detail=detail)
    return {"status": "ok"}

@app.post("/api/dossiers/update_coordonnees")
def modify_coordonnees(data: CoordonneesUpdateSchema):
    if not data.email.strip() and not data.telephone.strip():
        raise HTTPException(status_code=400, detail="Ajoute au moins un moyen de contact : e-mail ou téléphone.")
    ok, detail = update_coordonnees(data.client_id, data.email.strip(), data.telephone.strip())
    if not ok:
        raise HTTPException(status_code=502, detail=detail)
    return {"status": "ok"}

@app.delete("/api/dossiers/{client_id}")
def remove_dossier(client_id: int):
    ok, detail = delete_client(client_id)
    if not ok:
        raise HTTPException(status_code=502, detail=detail)
    return {"status": "ok"}

@app.post("/api/relancer")
def relancer(data: RelanceSchema):
    if not log_relance(data.client_id, data.email, data.pieces):
        raise HTTPException(status_code=502, detail="Impossible d'enregistrer la relance")
    return {"status": "ok"}

@app.get("/api/historique/{client_id}")
def fetch_historique(client_id: int):
    return get_historique(client_id)

@app.get("/api/export/excel")
def export_excel():
    workbook = Workbook()
    dossiers_sheet = workbook.active
    dossiers_sheet.title = "Dossiers"
    dossiers_sheet.append([
        "Nom", "Prénom", "E-mail", "Téléphone", "N° contrat", "Type de véhicule",
        "Date d'effet", "Marque", "Immatriculation", "Pièces manquantes", "Statut",
        "Dernière relance", "Commentaire"
    ])

    clients = get_all_clients()
    for client in clients:
        contrats = client.get("contrats", []) or [None]
        for contrat in contrats:
            pieces = (contrat or {}).get("pieces_manquantes", []) or []
            statut = client.get("texte_statut", "")
            dossiers_sheet.append([
                client.get("nom", ""),
                client.get("prenom", ""),
                client.get("email", ""),
                client.get("telephone", ""),
                (contrat or {}).get("num_contrat", ""),
                (contrat or {}).get("type_vehicule", ""),
                (contrat or {}).get("date_effet", ""),
                (contrat or {}).get("marque", ""),
                (contrat or {}).get("immat", ""),
                ", ".join(str(piece) for piece in pieces),
                statut,
                client.get("derniere_relance", "") or "",
                client.get("commentaire", "") or ""
            ])

    relances_sheet = workbook.create_sheet("Historique des relances")
    relances_sheet.append(["Nom", "Prénom", "Date et heure", "E-mail", "Pièces réclamées"])
    for client in clients:
        for relance in get_historique(client.get("id")):
            pieces = relance.get("pieces", []) or []
            relances_sheet.append([
                client.get("nom", ""),
                client.get("prenom", ""),
                relance.get("date_heure", ""),
                relance.get("email", ""),
                ", ".join(str(piece) for piece in pieces)
            ])

    entete_fill = PatternFill(fill_type="solid", fgColor="202124")
    entete_font = Font(color="FFCC00", bold=True)
    for sheet in (dossiers_sheet, relances_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = entete_fill
            cell.font = entete_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths_dossiers = [20, 20, 30, 18, 16, 20, 16, 18, 20, 48, 30, 24, 40]
    for index, width in enumerate(widths_dossiers, start=1):
        dossiers_sheet.column_dimensions[chr(64 + index)].width = width
    for column, width in {"A": 20, "B": 20, "C": 24, "D": 30, "E": 48}.items():
        relances_sheet.column_dimensions[column].width = width

    for row in dossiers_sheet.iter_rows(min_row=2):
        statut = str(row[10].value or "")
        if statut.startswith("Urgent"):
            row[10].fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
        elif "relance à faire" in statut.lower():
            row[10].fill = PatternFill(fill_type="solid", fgColor="FCE5CD")
        elif statut.startswith("Complet"):
            row[10].fill = PatternFill(fill_type="solid", fgColor="D9EAD3")

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = f"sauvegarde_relances_assurance_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
